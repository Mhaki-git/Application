"""
Ecrit les parametres saisis dans l'interface vers une COPIE du fichier Excel,
dans les memes cellules que celles lues par dataextraction.py. Ne touche
jamais aux donnees de consommation/PV (colonnes A-AC de DONNEES ENERGIE,
onglet PVS) : uniquement le bloc de parametres.
"""
import shutil
import openpyxl

# cle du dict de parametres -> adresse de cellule dans l'onglet "DONNEES ENERGIE"
PARAM_CELL_MAP = {
    "kwc": "F5",
    "kva_onduleur": "F6",
    "battery_power_kw": "F13",
    "battery_capacity_kwh": "F14",

    "old_cout_additionnel_contrat": "J6",
    "old_price_kwh_fixe": "J7",
    "old_price_inj_fixe": "J8",
    "price_hp": "J9",
    "price_hc": "J10",
    "price_inj_hp": "J11",
    "price_inj_hc": "J12",
    "prix_kwc_pv": "J13",
    "prix_kwh_batterie": "J14",
    "heure_debut_hp": "J15",
    "heure_debut_hc": "J16",
    "contrat_kw": "N16",

    "prix_vente_kwh": "J27",
    "prix_injection_fournisseur": "J28",
    "marge_fournisseur": "J29",
    "taxes_couts_proportionnels": "J30",
    "tarif_capacitaire_fournisseur": "J31",
    "marge_injection": "J32",

    "inflation_pct": "N12",
    "maintenance_eur_an": "N13",
    "degradation_pv_pct": "N14",
    "degradation_batterie_pct": "N15",
}


def write_params_to_copy(source_xlsm_path: str, dest_xlsm_path: str, params: dict) -> str:
    """
    Copie le fichier source vers dest_xlsm_path, puis ecrit les valeurs de
    `params` (sous-ensemble de PARAM_CELL_MAP) dans les cellules correspondantes
    de l'onglet DONNEES ENERGIE. Les formules, macros et donnees conso/PV du
    fichier original restent intactes.
    """
    shutil.copyfile(source_xlsm_path, dest_xlsm_path)

    wb = openpyxl.load_workbook(dest_xlsm_path, keep_vba=True)
    ws = wb["DONNEES ENERGIE"]

    for key, value in params.items():
        if key not in PARAM_CELL_MAP or value is None:
            continue
        ws[PARAM_CELL_MAP[key]] = value

    wb.save(dest_xlsm_path)
    wb.close()
    return dest_xlsm_path


def read_current_params(xlsm_path: str) -> dict:
    """Relit rapidement les valeurs actuelles (pour pre-remplir le formulaire)."""
    wb = openpyxl.load_workbook(xlsm_path, data_only=True, read_only=True)
    ws = wb["DONNEES ENERGIE"]
    current = {key: ws[cell].value for key, cell in PARAM_CELL_MAP.items()}
    wb.close()
    return current
