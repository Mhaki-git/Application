"""
Extraction DYNAMIQUE des donnees et parametres depuis le fichier Excel.

Principe : RIEN n'est fige en dur dans le code. A chaque execution, on relit
l'Excel tel qu'il est *maintenant*. Si tu changes une valeur dans le fichier
et relances le script, tout est recalcule automatiquement.
"""
import openpyxl
import pandas as pd
import numpy as np

XLSM_PATH = "ENERGY MIX ANALYSIS - AMELIORATION (2).xlsm"

MAX_SCAN_ROWS_QH = 40000
MAX_SCAN_ROWS_H = 8900

_CHAMPS_A_VERIFIER_SI_ZERO = {
    "maintenance_eur_an": "coût de maintenance annuel PV+batterie",
    "old_cout_additionnel_contrat": "frais fixes/abonnement de l'ancien contrat client",
    "tarif_capacitaire_fournisseur": "tarif capacitaire (peak shaving) -- si 0, le LP n'a "
                                      "aucune incitation à écrêter les pointes de puissance",
}


def _read_parameters(ws_de) -> dict:
    params = {
        "kwc": ws_de["F5"].value or 0,
        "kva_onduleur": ws_de["F6"].value or 0,
        "battery_power_kw": ws_de["F13"].value or 0,
        "battery_capacity_kwh": ws_de["F14"].value or 0,

        "old_cout_additionnel_contrat": ws_de["J6"].value or 0,
        "old_price_kwh_fixe": ws_de["J7"].value or 0,
        "old_price_inj_fixe": ws_de["J8"].value or 0,
        "price_hp": ws_de["J9"].value or 0,
        "price_hc": ws_de["J10"].value or 0,
        "price_inj_hp": ws_de["J11"].value or 0,
        "price_inj_hc": ws_de["J12"].value or 0,
        "prix_kwc_pv": ws_de["J13"].value or 0,
        "prix_kwh_batterie": ws_de["J14"].value or 0,
        "heure_debut_hp": ws_de["J15"].value,
        "heure_debut_hc": ws_de["J16"].value,
        "contrat_kw": ws_de["N16"].value or 0,

        "prix_vente_kwh": ws_de["J27"].value or 0,
        "prix_injection_fournisseur": ws_de["J28"].value or 0,
        "marge_fournisseur": ws_de["J29"].value or 0,
        "taxes_couts_proportionnels": ws_de["J30"].value or 0,
        "tarif_capacitaire_fournisseur": ws_de["J31"].value or 0,
        "marge_injection": ws_de["J32"].value or 0,

        "inflation_pct": ws_de["N12"].value or 0,
        "maintenance_eur_an": ws_de["N13"].value or 0,
        "degradation_pv_pct": ws_de["N14"].value or 0,
        "degradation_batterie_pct": ws_de["N15"].value or 0,
    }

    for champ, description in _CHAMPS_A_VERIFIER_SI_ZERO.items():
        if not params[champ]:
            print(f"  /!\\ {champ} = 0 ({description}) -- confirme que c'est bien voulu, "
                  f"sinon corrige la cellule Excel correspondante.")

    return params


def _read_conso_series(ws_de, max_scan_rows: int = MAX_SCAN_ROWS_QH) -> pd.Series:
    dates = []
    conso_kw = []
    for row in ws_de.iter_rows(min_row=2, max_row=1 + max_scan_rows,
                                min_col=1, max_col=29, values_only=True):
        d = row[0]
        if d is None:
            break
        dates.append(d)
        conso_kw.append(row[28] if row[28] is not None else 0.0)

    if not dates:
        raise ValueError("Aucune date valide trouvee en colonne A de DONNEES ENERGIE.")

    conso_kw = np.array(conso_kw, dtype=float)
    conso_kwh = conso_kw * 0.25
    series = pd.Series(conso_kwh, index=pd.DatetimeIndex(dates), name="conso_kwh")

    if series.index.duplicated().any():
        n_dup = int(series.index.duplicated().sum())
        print(f"  /!\\ {n_dup} timestamps dupliques dans la conso -- on garde la premiere occurrence.")
        series = series[~series.index.duplicated(keep="first")]

    return series


def _read_pv_series(ws_pv, kwc: float, max_scan_rows: int = MAX_SCAN_ROWS_H) -> pd.Series:
    dates = []
    pv_1kwc = []
    for row in ws_pv.iter_rows(min_row=5, max_row=4 + max_scan_rows,
                                min_col=1, max_col=116, values_only=True):
        d = row[0]
        if d is None:
            break
        dates.append(d)
        pv_1kwc.append(row[115] if row[115] is not None else 0.0)

    if not dates:
        raise ValueError("Aucune date valide trouvee en colonne A de PVS.")

    pv_1kwc = np.array(pv_1kwc, dtype=float)
    pv_kwh_per_quarter = pv_1kwc * kwc / 4.0

    qh_index = []
    qh_values = []
    for t, v in zip(dates, pv_kwh_per_quarter):
        for m in (0, 15, 30, 45):
            qh_index.append(t + pd.Timedelta(minutes=m))
            qh_values.append(v)

    series = pd.Series(qh_values, index=pd.DatetimeIndex(qh_index), name="pv_kwh")

    if series.index.duplicated().any():
        series = series[~series.index.duplicated(keep="first")]

    return series


def extract_all(xlsm_path: str = XLSM_PATH, kwc: float = None):
    wb = openpyxl.load_workbook(xlsm_path, data_only=True, read_only=True)
    ws_de = wb["DONNEES ENERGIE"]
    ws_pv = wb["PVS"]

    params = _read_parameters(ws_de)
    if kwc is None:
        kwc = params["kwc"]

    print("Lecture conso quart-horaire (arret dynamique en fin de donnees)...")
    conso_series = _read_conso_series(ws_de)
    print(f"  -> {len(conso_series)} points, de {conso_series.index[0]} a {conso_series.index[-1]}")

    print("Lecture profil PV horaire (arret dynamique en fin de donnees)...")
    pv_series = _read_pv_series(ws_pv, kwc)
    print(f"  -> {len(pv_series)} points (apres desagregation QH), "
          f"de {pv_series.index[0]} a {pv_series.index[-1]}")

    wb.close()

    df = pd.concat([conso_series, pv_series], axis=1).sort_index()
    n_missing_conso = int(df["conso_kwh"].isna().sum())
    n_missing_pv = int(df["pv_kwh"].isna().sum())
    if n_missing_conso or n_missing_pv:
        print(f"  /!\\ Alignement conso/PV : {n_missing_conso} pas sans conso, "
              f"{n_missing_pv} pas sans PV -- combles a 0.")
    df["conso_kwh"] = df["conso_kwh"].fillna(0.0)
    df["pv_kwh"] = df["pv_kwh"].fillna(0.0)

    df.index.name = "datetime"
    df["hour"] = df.index.hour
    df["is_weekend"] = df.index.dayofweek >= 5
    return params, df


def _parse_decimal_str(value) -> float:
    """Convertit une valeur (deja numerique, ou texte type '1234,56' /
    '1.234,56' / '1,234.56' / '1234.56') en float, en devinant le format
    decimal/milliers a partir du texte lui-meme -- ne suppose jamais un
    format fixe, car un export Excel FR utilise la virgule decimale alors
    qu'un export "standard" (ou un autre logiciel) utilise le point."""
    if value is None or value == "" or (isinstance(value, float) and pd.isna(value)):
        return float("nan")
    if isinstance(value, (int, float)):
        return float(value)

    s = str(value).strip().replace(" ", "").replace("\u00a0", "")  # espace normal + insecable
    if not s:
        return float("nan")

    if "," in s and "." in s:
        # Les deux presents : le DERNIER des deux est le separateur decimal,
        # l'autre est un separateur de milliers a supprimer.
        # ex: "1.234,56" (FR) -> 1234.56  |  "1,234.56" (US) -> 1234.56
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        # Une seule virgule -> on suppose le format europeen (decimale)
        s = s.replace(",", ".")
    # sinon : deja au format standard (point ou entier), rien a faire

    try:
        return float(s)
    except ValueError:
        return float("nan")


def _read_csv_robuste(csv_path_or_buffer) -> pd.DataFrame:
    """Lit un CSV en detectant automatiquement le separateur de colonnes
    (',' ou ';' -- un export Excel FR utilise generalement ';' puisque ','
    est deja pris par la virgule decimale). Le format decimal des valeurs
    est gere separement dans _parse_decimal_str, colonne par colonne."""
    raw = _read_raw_text(csv_path_or_buffer)
    first_line = raw.splitlines()[0] if raw.splitlines() else ""
    sep = ";" if first_line.count(";") >= first_line.count(",") else ","

    import io
    return pd.read_csv(io.StringIO(raw), sep=sep)


def _read_raw_text(csv_path_or_buffer) -> str:
    if hasattr(csv_path_or_buffer, "read"):
        raw = csv_path_or_buffer.read()
        if isinstance(raw, bytes):
            raw = raw.decode("utf-8-sig", errors="replace")  # utf-8-sig : gere aussi le BOM Excel
    else:
        with open(csv_path_or_buffer, "rb") as f:
            raw = f.read().decode("utf-8-sig", errors="replace")
    return raw


def _find_first_timestamp_row(lines, sep: str, max_scan: int = 200) -> int:
    """
    Cherche, parmi les premieres lignes du fichier, la premiere dont le champ
    de la 1ere colonne se lit comme une date/heure valide (ex: 01/01/2025,
    2025-01-01 00:00, etc). Ca permet de sauter automatiquement d'eventuelles
    lignes d'en-tete/metadonnees d'export (titres, noms de colonnes, unites,
    lignes vides...) sans avoir besoin de connaitre leur nombre a l'avance --
    on lit juste les donnees a partir de la 1ere date trouvee en colonne A.

    Retourne l'index (0-based, dans `lines`) de la premiere ligne valide. Si
    rien n'est trouve dans les `max_scan` premieres lignes, retourne 0 (aucun
    saut -- comportement inchange, la ligne 0 est traitee comme avant).
    """
    for i, line in enumerate(lines[:max_scan]):
        if not line.strip():
            continue
        first_field = line.split(sep)[0].strip().strip('"').strip("'")
        if not first_field:
            continue
        ts = pd.to_datetime(first_field, dayfirst=True, errors="coerce")
        if pd.notna(ts):
            return i
    return 0


def read_conso_csv(csv_path_or_buffer, timestamp_col: str = None,
                    value_col: str = None, unit: str = "kW") -> pd.Series:
    """
    Lit la consommation du client depuis un CSV simple, en remplacement de la
    lecture de la colonne AC de l'onglet DONNEES ENERGIE.

    Format attendu : 2 colonnes -- un timestamp et une valeur de consommation.
    Detection automatique des noms de colonnes si timestamp_col/value_col ne
    sont pas fournis (recherche insensible a la casse parmi des noms usuels).
    Detection automatique aussi du separateur de colonnes (',' ou ';') et du
    format decimal (point ou virgule) -- gere les exports Excel FR comme les
    exports "standard", sans that l'utilisateur ait a s'en soucier.

    unit : "kW" (puissance moyenne sur le pas de temps -- comme l'ancienne
           colonne AC) ou "kWh" (energie deja au pas de temps). Si "kW", la
           conversion en kWh se fait automatiquement a partir du pas de temps
           detecte dans les donnees (pas besoin que ce soit du quart-horaire :
           15 min, 30 min ou 60 min sont tous geres).

    Retourne une pd.Series indexee par datetime, en kWh par pas de temps,
    au MEME pas de temps que les donnees fournies (pas de reechantillonnage
    force en quart-horaire ici -- voir align_conso_to_quarter_hour).
    """
    df = _read_csv_robuste(csv_path_or_buffer)

    if timestamp_col is None:
        candidates = ["timestamp", "datetime", "date", "date/heure", "date_heure"]
        timestamp_col = next((c for c in df.columns if c.strip().lower() in candidates), df.columns[0])
    if value_col is None:
        candidates = ["conso", "consommation", "value", "valeur", "kw", "kwh", "power", "puissance"]
        value_col = next((c for c in df.columns if c.strip().lower() in candidates), df.columns[1])

    ts = pd.to_datetime(df[timestamp_col], dayfirst=True)
    values = df[value_col].apply(_parse_decimal_str)
    n_bad = int(values.isna().sum())
    if n_bad:
        print(f"  /!\\ {n_bad} valeurs de consommation illisibles dans le CSV -- remplacees par 0. "
              f"Verifie le format de la colonne '{value_col}' si ce nombre est eleve.")
    values = values.fillna(0.0)
    series = pd.Series(values.values, index=pd.DatetimeIndex(ts), name="conso_raw")
    series = series[~series.index.duplicated(keep="first")].sort_index()

    if len(series) < 2:
        raise ValueError("CSV de consommation : pas assez de lignes pour deduire le pas de temps.")

    step_minutes = int(round((series.index[1] - series.index[0]).total_seconds() / 60))
    if step_minutes <= 0:
        raise ValueError("Impossible de deduire un pas de temps positif depuis le CSV de consommation.")

    if unit.lower() == "kw":
        conso_kwh = series.values * (step_minutes / 60.0)
    elif unit.lower() == "kwh":
        conso_kwh = series.values
    else:
        raise ValueError(f"unit doit etre 'kW' ou 'kWh', recu : {unit!r}")

    print(f"  -> CSV conso : {len(series)} points, pas de temps detecte = {step_minutes} min, "
          f"total = {conso_kwh.sum():,.0f} kWh")

    return pd.Series(conso_kwh, index=series.index, name="conso_kwh")


def align_conso_to_quarter_hour(conso_kwh: pd.Series) -> pd.Series:
    """
    Ramene une serie de conso (kWh par pas de temps, pas quelconque) a une
    grille quart-horaire reguliere, comme l'attend le moteur LP (DT=0.25).

    - Si le pas source est deja 15 min : ne change rien.
    - Si le pas source est plus grossier (30 min, 60 min) : repartit l'energie
      a parts egales sur les quarts d'heure inclus (pas d'invention de forme
      de courbe -- simple etalement uniforme, a documenter dans le rapport).
    - Si le pas source est plus fin que 15 min : agrege par sommation.
    """
    step_minutes = int(round((conso_kwh.index[1] - conso_kwh.index[0]).total_seconds() / 60))
    if step_minutes == 15:
        return conso_kwh

    if step_minutes > 15:
        n_qh = step_minutes // 15
        qh_index = []
        qh_values = []
        for ts, val in conso_kwh.items():
            share = val / n_qh
            for m in range(n_qh):
                qh_index.append(ts + pd.Timedelta(minutes=15 * m))
                qh_values.append(share)
        result = pd.Series(qh_values, index=pd.DatetimeIndex(qh_index), name="conso_kwh")
    else:
        result = conso_kwh.resample("15min").sum()

    print(f"  /!\\ Conso reechantillonnee de {step_minutes} min vers 15 min "
          f"({'repartition uniforme' if step_minutes > 15 else 'agregation'}).")
    return result


def build_timeseries_from_sources(conso_csv_path_or_buffer, pv_hourly_profile_1kwc: pd.Series,
                                    kwc: float, unit: str = "kW") -> pd.DataFrame:
    """
    Construit le DataFrame (conso_kwh, pv_kwh, hour, is_weekend) attendu par
    le moteur de dispatch (fournisseur_roi.py), a partir :
    - d'un CSV de consommation client (voir read_conso_csv)
    - d'un profil PV PVGIS pour 1 kWc (voir pv_pvgis.fetch_pv_profile_1kwc)

    C'est le remplacement direct de extract_all() pour la nouvelle interface
    sans Excel : le reste du pipeline (run_fournisseur_model, etc.) n'a besoin
    d'aucune modification, il consomme ce DataFrame de la meme facon.
    """
    from pv_pvgis import expand_to_quarter_hour

    print("Lecture de la consommation client (CSV)...")
    conso_raw = read_conso_csv(conso_csv_path_or_buffer, unit=unit)
    conso_qh = align_conso_to_quarter_hour(conso_raw)

    print("Alignement du profil PV (PVGIS) sur la grille quart-horaire du client...")
    pv_qh = expand_to_quarter_hour(pv_hourly_profile_1kwc, kwc=kwc, target_index=conso_qh.index)

    df = pd.concat([conso_qh, pv_qh], axis=1).sort_index()
    df["conso_kwh"] = df["conso_kwh"].fillna(0.0)
    df["pv_kwh"] = df["pv_kwh"].fillna(0.0)
    df.index.name = "datetime"
    df["hour"] = df.index.hour
    df["is_weekend"] = df.index.dayofweek >= 5

    print(f"  -> Serie finale : {len(df)} points, conso={df['conso_kwh'].sum():,.0f} kWh/an, "
          f"pv={df['pv_kwh'].sum():,.0f} kWh/an")
    return df


def extract_parameters(xlsm_path: str = XLSM_PATH) -> dict:
    params, _ = extract_all(xlsm_path)
    return params


def extract_timeseries(xlsm_path: str = XLSM_PATH, kwc: float = None) -> pd.DataFrame:
    _, df = extract_all(xlsm_path, kwc=kwc)
    return df
