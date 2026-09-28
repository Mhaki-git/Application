"""
Ecrit les parametres saisis dans l'interface vers une COPIE du fichier Excel,
dans les memes cellules que celles lues par dataextraction.py. Ne touche
jamais aux donnees de consommation/PV (colonnes A-AC de DONNEES ENERGIE,
onglet PVS) : uniquement le bloc de parametres.

STATUT : orphelin fonctionnel -- aucun appelant dans app.py ni engine/ a ce
jour (l'UI Streamlit actuelle ne relit/re-ecrit jamais le .xlsm). Conserve
comme outil disponible pour un usage manuel/futur (regenerer une copie
parametree du classeur source), pas comme partie active du pipeline.
"""
import shutil
import openpyxl

from io_excel.excel_schema import PARAM_CELL_MAP, PARAMETERS_SHEET_NAME


def write_params_to_copy(source_xlsm_path: str, dest_xlsm_path: str, params: dict) -> str:
    """
    Copie le fichier source vers dest_xlsm_path, puis ecrit les valeurs de
    `params` (sous-ensemble de PARAM_CELL_MAP) dans les cellules correspondantes
    de l'onglet DONNEES ENERGIE. Les formules, macros et donnees conso/PV du
    fichier original restent intactes.
    """
    # La copie se fait AVANT ouverture avec openpyxl : le fichier source original
    # n'est jamais ouvert en ecriture, seule la copie dest_xlsm_path est modifiee.
    shutil.copyfile(source_xlsm_path, dest_xlsm_path)

    # keep_vba=True est indispensable pour un .xlsm : sans cette option, openpyxl
    # supprimerait les macros VBA du classeur a la sauvegarde.
    wb = openpyxl.load_workbook(dest_xlsm_path, keep_vba=True)
    ws = wb[PARAMETERS_SHEET_NAME]

    for key, value in params.items():
        # Ignore silencieusement les cles absentes du mapping (parametres non geres
        # par ce schema) et les valeurs None (champ non renseigne) : on ne veut pas
        # ecraser une cellule existante avec une valeur vide.
        if key not in PARAM_CELL_MAP or value is None:
            continue
        ws[PARAM_CELL_MAP[key]] = value

    wb.save(dest_xlsm_path)
    wb.close()
    return dest_xlsm_path


def read_current_params(xlsm_path: str) -> dict:
    """Relit rapidement les valeurs actuelles (pour pre-remplir le formulaire)."""
    # data_only=True renvoie la derniere valeur calculee des formules (pas la formule
    # elle-meme) ; read_only=True accelere la lecture sur un gros classeur car il
    # n'ouvre pas la feuille entiere en memoire pour modification.
    wb = openpyxl.load_workbook(xlsm_path, data_only=True, read_only=True)
    ws = wb[PARAMETERS_SHEET_NAME]
    current = {key: ws[cell].value for key, cell in PARAM_CELL_MAP.items()}
    wb.close()
    return current
