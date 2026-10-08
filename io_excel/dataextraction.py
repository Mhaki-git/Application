"""
Extraction DYNAMIQUE des donnees et parametres depuis le fichier Excel.

Principe : RIEN n'est fige en dur dans le code. A chaque execution, on relit
l'Excel tel qu'il est *maintenant*. Si tu changes une valeur dans le fichier
et relances le script, tout est recalcule automatiquement.

STATUT : `extract_all` (et ses dependances `_read_parameters`,
`_read_conso_series`, `_read_pv_series`) constituent le chemin "legacy Excel"
-- lecture complete d'un .xlsm (parametres + series temporelles). L'UI
Streamlit actuelle (app.py) NE PASSE PLUS par ce chemin : elle utilise
`build_timeseries_from_sources` (CSV/profil integre + PVGIS/PV fige) et
passe les parametres directement en dict. `extract_all` reste utilise par
`engine.fournisseur_roi.run_fournisseur_model` (point d'entree CLI/scripts,
non appele par app.py) et par les tests de non-regression. Conserve pour
compatibilite descendante, pas le flux de production.
"""
import openpyxl
import pandas as pd
import numpy as np

from io_excel.excel_schema import PARAM_CELL_MAP, PARAMETERS_SHEET_NAME
from engine.calendar_utils import NearestByCalendarSymmetry

XLSM_PATH = "ENERGY MIX ANALYSIS - AMELIORATION (2).xlsm"

# Cles dont la valeur brute (y compris None) doit etre conservee telle quelle,
# sans retomber a 0 si la cellule Excel est vide -- toutes les autres cles de
# PARAM_CELL_MAP utilisent `valeur or 0`.
_KEYS_WITHOUT_ZERO_FALLBACK = {"heure_debut_hp", "heure_debut_hc"}

# Limites de securite pour le parcours ligne par ligne des onglets Excel
# (arret normal = premiere colonne A vide, ces bornes ne servent qu'a eviter
# un parcours illimite sur un fichier corrompu). 40000 lignes ~ un an de
# quart-heures avec large marge ; 8900 lignes ~ un an d'heures avec marge.
MAX_SCAN_ROWS_QH = 40000
MAX_SCAN_ROWS_H = 8900

_CHAMPS_A_VERIFIER_SI_ZERO = {
    "maintenance_eur_an": "coût de maintenance annuel PV+batterie",
    "old_cout_additionnel_contrat": "frais fixes/abonnement de l'ancien contrat client",
    "tarif_capacitaire_fournisseur": "tarif capacitaire (peak shaving) -- si 0, le LP n'a "
                                      "aucune incitation à écrêter les pointes de puissance",
}


def _read_parameters(ws_de) -> dict:
    """
    Lit tous les parametres scalaires de l'onglet DONNEES ENERGIE, a partir des
    adresses de cellules definies dans PARAM_CELL_MAP (source de verite partagee
    avec excel_writer.py).

    Piege : une cellule vide donne `None` en openpyxl. Pour la plupart des cles,
    on retombe sur 0 (`valeur or 0`) afin d'eviter des `None` qui casseraient les
    calculs downstream. Exception : les cles listees dans
    _KEYS_WITHOUT_ZERO_FALLBACK (heures de debut HP/HC) ou `None` a un sens
    different de 0 (0 serait interprete comme "minuit"), donc la valeur brute
    est conservee telle quelle.

    Retourne un dict {cle_parametre: valeur}, tel que defini par PARAM_CELL_MAP.
    """
    params = {}
    for key, cell in PARAM_CELL_MAP.items():
        value = ws_de[cell].value
        params[key] = value if key in _KEYS_WITHOUT_ZERO_FALLBACK else (value or 0)

    # Garde-fou : certains parametres a 0 sont techniquement valides mais
    # neutralisent une partie du modele economique sans erreur explicite --
    # on avertit au lieu d'echouer, pour laisser la main a l'utilisateur.
    for champ, description in _CHAMPS_A_VERIFIER_SI_ZERO.items():
        if not params[champ]:
            print(f"  /!\\ {champ} = 0 ({description}) -- confirme que c'est bien voulu, "
                  f"sinon corrige la cellule Excel correspondante.")

    return params


def _read_conso_series(ws_de, max_scan_rows: int = MAX_SCAN_ROWS_QH) -> pd.Series:
    """
    Lit la serie de consommation quart-horaire du client depuis l'onglet
    DONNEES ENERGIE : colonne A = timestamp, colonne AC (index 28, 0-based)
    = puissance moyenne en kW sur le quart d'heure.

    L'arret de lecture est dynamique : on scanne ligne par ligne jusqu'a
    trouver une colonne A vide (fin des donnees), dans la limite de
    max_scan_rows pour eviter un parcours illimite si le fichier est corrompu.

    Piege d'unite : la colonne AC est une PUISSANCE en kW (pas une energie),
    d'ou la conversion `* 0.25` pour obtenir l'energie en kWh sur le quart
    d'heure (0.25 h). Les cellules vides sont traitees comme 0 kW.

    Retourne une pd.Series en kWh par quart d'heure, indexee par datetime,
    dedupliquee (on garde la premiere occurrence en cas de timestamp repete).
    """
    dates = []
    conso_kw = []
    for row in ws_de.iter_rows(min_row=2, max_row=1 + max_scan_rows,
                                min_col=1, max_col=29, values_only=True):
        d = row[0]
        if d is None:
            break
        dates.append(d)
        conso_kw.append(row[28] if row[28] is not None else 0.0)

    if not dates:
        raise ValueError("Aucune date valide trouvee en colonne A de DONNEES ENERGIE.")

    conso_kw = np.array(conso_kw, dtype=float)
    conso_kwh = conso_kw * 0.25  # kW moyen sur 15 min -> kWh (0.25 h)
    series = pd.Series(conso_kwh, index=pd.DatetimeIndex(dates), name="conso_kwh")

    if series.index.duplicated().any():
        n_dup = int(series.index.duplicated().sum())
        print(f"  /!\\ {n_dup} timestamps dupliques dans la conso -- on garde la premiere occurrence.")
        series = series[~series.index.duplicated(keep="first")]

    return series


def _read_pv_series(ws_pv, kwc: float, max_scan_rows: int = MAX_SCAN_ROWS_H) -> pd.Series:
    """
    Lit le profil de production PV depuis l'onglet PVS : colonne A = timestamp
    (au pas HORAIRE, contrairement a la conso qui est quart-horaire), colonne
    DL (index 115, 0-based) = production normalisee pour une installation de
    1 kWc, en kWh par heure.

    Les donnees demarrent a la ligne 5 (lignes 1-4 = en-tetes/metadonnees de
    l'export PVS). Arret dynamique de lecture des que la colonne A est vide,
    dans la limite de max_scan_rows.

    Mise a l'echelle : la production pour 1 kWc est multipliee par `kwc`
    (puissance crete reelle de l'installation, en kWc) pour obtenir l'energie
    horaire reelle, puis divisee par 4 pour repartir uniformement cette
    energie horaire sur les 4 quarts d'heure de l'heure (desagregation
    horaire -> quart-horaire par simple division, sans modeliser de variation
    infra-horaire).

    Retourne une pd.Series en kWh par quart d'heure (index quart-horaire),
    alignee sur la meme granularite que la conso, dedupliquee.
    """
    dates = []
    pv_1kwc = []
    for row in ws_pv.iter_rows(min_row=5, max_row=4 + max_scan_rows,
                                min_col=1, max_col=116, values_only=True):
        d = row[0]
        if d is None:
            break
        dates.append(d)
        pv_1kwc.append(row[115] if row[115] is not None else 0.0)

    if not dates:
        raise ValueError("Aucune date valide trouvee en colonne A de PVS.")

    pv_1kwc = np.array(pv_1kwc, dtype=float)
    pv_kwh_per_quarter = pv_1kwc * kwc / 4.0  # kWh/h pour l'installation -> kWh par quart d'heure

    # Desagregation horaire -> quart-horaire : on duplique la meme valeur
    # (energie/4) sur les 4 sous-pas de 15 min de chaque heure source.
    qh_index = []
    qh_values = []
    for t, v in zip(dates, pv_kwh_per_quarter):
        for m in (0, 15, 30, 45):
            qh_index.append(t + pd.Timedelta(minutes=m))
            qh_values.append(v)

    series = pd.Series(qh_values, index=pd.DatetimeIndex(qh_index), name="pv_kwh")

    if series.index.duplicated().any():
        series = series[~series.index.duplicated(keep="first")]

    return series


def extract_all(xlsm_path: str = XLSM_PATH, kwc: float = None):
    """
    Point d'entree "legacy Excel" : lit integralement le classeur .xlsm
    (parametres + series temporelles conso/PV) et retourne (params, df) prets
    a l'emploi pour le moteur de dispatch. Voir le docstring de module pour le
    statut de cette fonction (chemin legacy, plus utilise par l'UI Streamlit).

    - `data_only=True` : recupere les valeurs calculees des formules Excel
      (pas les formules elles-memes).
    - `read_only=True` : mode lecture seule d'openpyxl, plus rapide et plus
      leger en memoire sur un gros classeur.

    Retourne (params: dict, df: pd.DataFrame) -- df contient conso_kwh,
    pv_kwh, hour, is_weekend, indexe par datetime quart-horaire.
    """
    wb = openpyxl.load_workbook(xlsm_path, data_only=True, read_only=True)
    ws_de = wb[PARAMETERS_SHEET_NAME]
    ws_pv = wb["PVS"]

    params = _read_parameters(ws_de)
    if kwc is None:
        kwc = params["kwc"]  # a defaut de kwc explicite, on utilise celui saisi dans le classeur

    print("Lecture conso quart-horaire (arret dynamique en fin de donnees)...")
    conso_series = _read_conso_series(ws_de)
    print(f"  -> {len(conso_series)} points, de {conso_series.index[0]} a {conso_series.index[-1]}")

    print("Lecture profil PV horaire (arret dynamique en fin de donnees)...")
    pv_series = _read_pv_series(ws_pv, kwc)
    print(f"  -> {len(pv_series)} points (apres desagregation QH), "
          f"de {pv_series.index[0]} a {pv_series.index[-1]}")

    wb.close()  # libere le classeur des l'extraction terminee (mode read_only garde le fichier ouvert sinon)

    # Jointure par index datetime : conso et PV n'ont pas forcement exactement
    # les memes timestamps (sources differentes, arrets de lecture independants) ;
    # le concat aligne sur l'union des index et laisse des NaN la ou une des deux
    # series n'a pas de valeur a ce pas de temps.
    df = pd.concat([conso_series, pv_series], axis=1).sort_index()
    n_missing_conso = int(df["conso_kwh"].isna().sum())
    n_missing_pv = int(df["pv_kwh"].isna().sum())
    if n_missing_conso or n_missing_pv:
        print(f"  /!\\ Alignement conso/PV : {n_missing_conso} pas sans conso, "
              f"{n_missing_pv} pas sans PV -- combles a 0.")
    df["conso_kwh"] = df["conso_kwh"].fillna(0.0)
    df["pv_kwh"] = df["pv_kwh"].fillna(0.0)

    df.index.name = "datetime"
    df["hour"] = df.index.hour
    df["is_weekend"] = df.index.dayofweek >= 5  # 5=samedi, 6=dimanche (convention pandas dayofweek)
    return params, df


def _parse_decimal_str(value) -> float:
    """Convertit une valeur (deja numerique, ou texte type '1234,56' /
    '1.234,56' / '1,234.56' / '1234.56') en float, en devinant le format
    decimal/milliers a partir du texte lui-meme -- ne suppose jamais un
    format fixe, car un export Excel FR utilise la virgule decimale alors
    qu'un export "standard" (ou un autre logiciel) utilise le point."""
    if value is None or value == "" or (isinstance(value, float) and pd.isna(value)):
        return float("nan")
    if isinstance(value, (int, float)):
        return float(value)

    s = str(value).strip().replace(" ", "").replace("\u00a0", "")  # espace normal + insecable
    if not s:
        return float("nan")

    if "," in s and "." in s:
        # Les deux presents : le DERNIER des deux est le separateur decimal,
        # l'autre est un separateur de milliers a supprimer.
        # ex: "1.234,56" (FR) -> 1234.56  |  "1,234.56" (US) -> 1234.56
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        # Une seule virgule -> on suppose le format europeen (decimale)
        s = s.replace(",", ".")
    # sinon : deja au format standard (point ou entier), rien a faire

    try:
        return float(s)
    except ValueError:
        return float("nan")


def _read_csv_robuste(csv_path_or_buffer) -> pd.DataFrame:
    """Lit un CSV en detectant automatiquement le separateur de colonnes
    (',' ou ';' -- un export Excel FR utilise generalement ';' puisque ','
    est deja pris par la virgule decimale). Le format decimal des valeurs
    est gere separement dans _parse_decimal_str, colonne par colonne."""
    raw = _read_raw_text(csv_path_or_buffer)
    first_line = raw.splitlines()[0] if raw.splitlines() else ""
    sep = ";" if first_line.count(";") >= first_line.count(",") else ","

    import io
    return pd.read_csv(io.StringIO(raw), sep=sep)


def _read_raw_text(csv_path_or_buffer) -> str:
    """
    Lit le contenu texte brut d'un CSV, qu'il soit fourni par chemin (str) ou
    par un objet fichier/buffer deja ouvert (ex: upload Streamlit, qui expose
    `.read()`).

    Le decodage force en "utf-8-sig" gere a la fois l'UTF-8 standard et le BOM
    (byte order mark) qu'Excel ajoute systematiquement en tete des CSV exportes
    depuis Windows -- sans ce mode, le BOM se retrouverait concatene au nom de
    la premiere colonne et casserait sa detection automatique en aval.
    `errors="replace"` evite un crash sur un octet mal encode isole plutot que
    de faire echouer tout le chargement.
    """
    if hasattr(csv_path_or_buffer, "read"):
        raw = csv_path_or_buffer.read()
        if isinstance(raw, bytes):
            raw = raw.decode("utf-8-sig", errors="replace")  # utf-8-sig : gere aussi le BOM Excel
    else:
        with open(csv_path_or_buffer, "rb") as f:
            raw = f.read().decode("utf-8-sig", errors="replace")
    return raw


def _find_first_timestamp_row(lines, sep: str, max_scan: int = 200) -> int:
    """
    Cherche, parmi les premieres lignes du fichier, la premiere dont le champ
    de la 1ere colonne se lit comme une date/heure valide (ex: 01/01/2025,
    2025-01-01 00:00, etc). Ca permet de sauter automatiquement d'eventuelles
    lignes d'en-tete/metadonnees d'export (titres, noms de colonnes, unites,
    lignes vides...) sans avoir besoin de connaitre leur nombre a l'avance --
    on lit juste les donnees a partir de la 1ere date trouvee en colonne A.

    Retourne l'index (0-based, dans `lines`) de la premiere ligne valide. Si
    rien n'est trouve dans les `max_scan` premieres lignes, retourne 0 (aucun
    saut -- comportement inchange, la ligne 0 est traitee comme avant).
    """
    for i, line in enumerate(lines[:max_scan]):
        if not line.strip():
            continue
        first_field = line.split(sep)[0].strip().strip('"').strip("'")
        if not first_field:
            continue
        ts = pd.to_datetime(first_field, dayfirst=True, errors="coerce")
        if pd.notna(ts):
            return i
    return 0


def _parse_timestamp_column(col: pd.Series) -> pd.DatetimeIndex:
    """
    Parse une colonne de timestamps en gerant a la fois le format ISO 8601
    (AAAA-MM-JJ, non ambigu) et la convention FR (JJ/MM/AAAA). Essaie
    d'abord ISO8601 strict : si une seule valeur n'y correspond pas, bascule
    sur dayfirst=True (JJ/MM/AAAA) pour toute la colonne.

    Necessaire car pd.to_datetime(col, dayfirst=True) seul, applique a des
    timestamps deja ISO (AAAA-MM-JJ), peut inferer le mauvais format a partir
    des premieres lignes (AAAA-JJ-MM) et planter des qu'une ligne suivante a
    un jour > 12 (ex: "2024-01-13" rejete car interprete comme "annee-jour-mois").
    """
    try:
        return pd.to_datetime(col, format="ISO8601")
    except (ValueError, TypeError):
        return pd.to_datetime(col, dayfirst=True, format="mixed")


_TIMESTAMP_NAMES = ["timestamp", "datetime", "date", "date/heure", "date_heure", "horodatage"]
_CONSO_NAMES = ["conso", "consommation", "value", "valeur", "kw", "kwh", "power", "puissance",
                "prelevement", "prélèvement", "offtake", "volume_prelevement"]
_INJECTION_NAMES = ["injection", "injections", "inj", "export", "exportation", "injecte", "injecté",
                    "production", "volume_injection"]


def list_csv_columns(csv_path_or_buffer) -> list:
    """Noms des colonnes d'un CSV (separateur ',' ou ';' detecte automatiquement),
    pour laisser l'utilisateur choisir horodatage / consommation / injection."""
    df = _read_csv_robuste(csv_path_or_buffer)
    return [str(c) for c in df.columns]


def guess_csv_columns(columns) -> dict:
    """Devine {timestamp, conso, injection} parmi les noms de colonnes (insensible
    a la casse, correspondance exacte puis par mot contenu). `injection` vaut None
    si aucune colonne ne ressemble a de l'injection."""
    cols = list(columns)
    low = {c: c.strip().lower() for c in cols}

    def pick(names, exclude=()):
        for c in cols:
            if c not in exclude and low[c] in names:
                return c
        for c in cols:
            if c not in exclude and any(n in low[c] for n in names if len(n) >= 3):
                return c
        return None

    ts = pick(_TIMESTAMP_NAMES) or cols[0]
    inj = pick(_INJECTION_NAMES, exclude=(ts,))
    conso = pick(_CONSO_NAMES, exclude=(ts, inj)) or next((c for c in cols if c not in (ts, inj)), cols[0])
    return {"timestamp": ts, "conso": conso, "injection": inj}


def read_conso_injection_csv(csv_path_or_buffer, timestamp_col: str = None, value_col: str = None,
                             injection_col: str = None, unit: str = "kW") -> tuple:
    """
    Comme read_conso_csv, mais lit en plus (optionnellement) une colonne
    d'INJECTION mesuree. Retourne (conso_kwh, injection_kwh) : deux pd.Series
    sur le meme index et au meme pas de temps natif, en kWh par pas ;
    injection_kwh vaut None si injection_col n'est pas fourni.
    Valeurs illisibles -> 0 ; les injections negatives sont ramenees a 0.
    """
    df = _read_csv_robuste(csv_path_or_buffer)
    guessed = guess_csv_columns(df.columns)
    timestamp_col = timestamp_col or guessed["timestamp"]
    value_col = value_col or guessed["conso"]

    ts = _parse_timestamp_column(df[timestamp_col])
    frame = pd.DataFrame({"conso_raw": df[value_col].apply(_parse_decimal_str).values}, index=pd.DatetimeIndex(ts))
    if injection_col:
        frame["injection_raw"] = df[injection_col].apply(_parse_decimal_str).values
    n_bad = int(frame.isna().sum().sum())
    if n_bad:
        print(f"  /!\\ {n_bad} valeurs illisibles dans le CSV -- remplacees par 0.")
    frame = frame.fillna(0.0)
    frame = frame[~frame.index.duplicated(keep="first")].sort_index()

    if len(frame) < 2:
        raise ValueError("CSV de consommation : pas assez de lignes pour deduire le pas de temps.")
    step_minutes = int(round((frame.index[1] - frame.index[0]).total_seconds() / 60))
    if step_minutes <= 0:
        raise ValueError("Impossible de deduire un pas de temps positif depuis le CSV de consommation.")

    if unit.lower() == "kw":
        factor = step_minutes / 60.0
    elif unit.lower() == "kwh":
        factor = 1.0
    else:
        raise ValueError(f"unit doit etre 'kW' ou 'kWh', recu : {unit!r}")

    conso = pd.Series(frame["conso_raw"].values * factor, index=frame.index, name="conso_kwh")
    injection = None
    if injection_col:
        injection = pd.Series(np.clip(frame["injection_raw"].values, 0.0, None) * factor,
                              index=frame.index, name="injection_mesuree_kwh")
    print(f"  -> CSV conso : {len(frame)} points, pas de temps = {step_minutes} min, "
          f"conso = {conso.sum():,.0f} kWh"
          + (f", injection = {injection.sum():,.0f} kWh" if injection is not None else ""))
    return conso, injection


def read_conso_csv(csv_path_or_buffer, timestamp_col: str = None,
                    value_col: str = None, unit: str = "kW") -> pd.Series:
    """
    Lit la consommation du client depuis un CSV simple, en remplacement de la
    lecture de la colonne AC de l'onglet DONNEES ENERGIE.

    Format attendu : 2 colonnes -- un timestamp et une valeur de consommation.
    Detection automatique des noms de colonnes si timestamp_col/value_col ne
    sont pas fournis (recherche insensible a la casse parmi des noms usuels).
    Detection automatique aussi du separateur de colonnes (',' ou ';') et du
    format decimal (point ou virgule) -- gere les exports Excel FR comme les
    exports "standard", sans that l'utilisateur ait a s'en soucier.

    unit : "kW" (puissance moyenne sur le pas de temps -- comme l'ancienne
           colonne AC) ou "kWh" (energie deja au pas de temps). Si "kW", la
           conversion en kWh se fait automatiquement a partir du pas de temps
           detecte dans les donnees (pas besoin que ce soit du quart-horaire :
           15 min, 30 min ou 60 min sont tous geres).

    Retourne une pd.Series indexee par datetime, en kWh par pas de temps,
    au MEME pas de temps que les donnees fournies (pas de reechantillonnage
    force en quart-horaire ici -- voir align_conso_to_quarter_hour).
    """
    df = _read_csv_robuste(csv_path_or_buffer)

    if timestamp_col is None:
        candidates = ["timestamp", "datetime", "date", "date/heure", "date_heure"]
        timestamp_col = next((c for c in df.columns if c.strip().lower() in candidates), df.columns[0])
    if value_col is None:
        candidates = ["conso", "consommation", "value", "valeur", "kw", "kwh", "power", "puissance"]
        value_col = next((c for c in df.columns if c.strip().lower() in candidates), df.columns[1])

    ts = _parse_timestamp_column(df[timestamp_col])
    values = df[value_col].apply(_parse_decimal_str)
    n_bad = int(values.isna().sum())
    if n_bad:
        print(f"  /!\\ {n_bad} valeurs de consommation illisibles dans le CSV -- remplacees par 0. "
              f"Verifie le format de la colonne '{value_col}' si ce nombre est eleve.")
    values = values.fillna(0.0)
    series = pd.Series(values.values, index=pd.DatetimeIndex(ts), name="conso_raw")
    series = series[~series.index.duplicated(keep="first")].sort_index()

    if len(series) < 2:
        raise ValueError("CSV de consommation : pas assez de lignes pour deduire le pas de temps.")

    # Le pas de temps n'est jamais suppose a l'avance : on le deduit de l'ecart
    # entre les deux premiers timestamps tries, ce qui permet de gerer 15/30/60 min
    # (ou tout autre pas regulier) sans configuration manuelle.
    step_minutes = int(round((series.index[1] - series.index[0]).total_seconds() / 60))
    if step_minutes <= 0:
        raise ValueError("Impossible de deduire un pas de temps positif depuis le CSV de consommation.")

    if unit.lower() == "kw":
        conso_kwh = series.values * (step_minutes / 60.0)  # puissance moyenne (kW) x duree (h) = energie (kWh)
    elif unit.lower() == "kwh":
        conso_kwh = series.values
    else:
        raise ValueError(f"unit doit etre 'kW' ou 'kWh', recu : {unit!r}")

    print(f"  -> CSV conso : {len(series)} points, pas de temps detecte = {step_minutes} min, "
          f"total = {conso_kwh.sum():,.0f} kWh")

    return pd.Series(conso_kwh, index=series.index, name="conso_kwh")


def read_dayahead_csv(csv_path_or_buffer, timestamp_col: str = None,
                       value_col: str = None) -> pd.Series:
    """
    Lit une serie de prix Day-Ahead (Belpex) depuis un CSV simple, en
    remplacement d'un upload .pkl -- pd.read_pickle sur un fichier fourni par
    l'utilisateur peut executer du code Python arbitraire (deserialisation
    non fiable), un CSV ne presente pas ce risque.

    Format attendu : 2 colonnes -- un timestamp et un prix en EUR/kWh.
    Detection automatique des noms de colonnes, du separateur (',' ou ';')
    et du format decimal (point ou virgule), meme logique que read_conso_csv.

    Retourne une pd.Series indexee par datetime (naive), prix en EUR/kWh, au
    pas de temps natif du CSV (le calage sur les dates client se fait plus
    loin par _align_market_price_to_calendar, quel que soit ce pas).
    """
    df = _read_csv_robuste(csv_path_or_buffer)

    if timestamp_col is None:
        candidates = ["timestamp", "datetime", "date", "date/heure", "date_heure"]
        timestamp_col = next((c for c in df.columns if c.strip().lower() in candidates), df.columns[0])
    if value_col is None:
        candidates = ["prix", "price", "value", "valeur", "eur_kwh", "eur/kwh", "prix_eur_kwh"]
        value_col = next((c for c in df.columns if c.strip().lower() in candidates), df.columns[1])

    ts = _parse_timestamp_column(df[timestamp_col])
    values = df[value_col].apply(_parse_decimal_str)
    n_bad = int(values.isna().sum())
    if n_bad:
        print(f"  /!\\ {n_bad} prix Day-Ahead illisibles dans le CSV -- ignores. "
              f"Verifie le format de la colonne '{value_col}' si ce nombre est eleve.")

    series = pd.Series(values.values, index=pd.DatetimeIndex(ts), name="price_eur_kwh")
    series = series.dropna()  # lignes illisibles (n_bad ci-dessus) retirees plutot que mises a 0 : un prix a 0 fausserait le ROI, contrairement a une conso a 0 qui est plausible
    series = series[~series.index.duplicated(keep="first")].sort_index()

    if len(series) < 2:
        raise ValueError("CSV Day-Ahead : pas assez de lignes valides pour construire la serie.")

    # Garde-fou d'unite : les prix Day-Ahead Belpex sont generalement publies en
    # EUR/MWh alors que le moteur attend des EUR/kWh -- une mediane superieure a
    # 2.0 EUR/kWh est trop elevee pour un prix electrique plausible et trahit
    # presque toujours un oubli de conversion (diviser par 1000) en amont.
    med = float(np.nanmedian(np.abs(series.values)))
    if med > 2.0:
        print(f"  /!\\ ATTENTION UNITES : la mediane des prix chargee est {med:.1f} -- "
              f"ca ressemble a des EUR/MWh, pas des EUR/kWh (colonne '{value_col}').")

    print(f"  -> CSV Day-Ahead : {len(series)} points, de {series.index[0]} a {series.index[-1]}")
    return series


def align_conso_to_quarter_hour(conso_kwh: pd.Series) -> pd.Series:
    """
    Ramene une serie de conso (kWh par pas de temps, pas quelconque) a une
    grille quart-horaire reguliere, comme l'attend le moteur LP (DT=0.25).

    - Si le pas source est deja 15 min : ne change rien.
    - Si le pas source est plus grossier (30 min, 60 min) : repartit l'energie
      a parts egales sur les quarts d'heure inclus (pas d'invention de forme
      de courbe -- simple etalement uniforme, a documenter dans le rapport).
    - Si le pas source est plus fin que 15 min : agrege par sommation.
    """
    step_minutes = int(round((conso_kwh.index[1] - conso_kwh.index[0]).total_seconds() / 60))
    if step_minutes == 15:
        return conso_kwh

    if step_minutes > 15:
        # Pas plus grossier que 15 min (ex: 30 ou 60 min) : on repartit l'energie
        # du pas source a parts EGALES sur les n_qh quarts d'heure qu'il couvre.
        # C'est un etalement uniforme, pas une reconstitution de la vraie courbe
        # de charge infra-pas -- l'information de forme est perdue a la source.
        n_qh = step_minutes // 15
        qh_index = []
        qh_values = []
        for ts, val in conso_kwh.items():
            share = val / n_qh
            for m in range(n_qh):
                qh_index.append(ts + pd.Timedelta(minutes=15 * m))
                qh_values.append(share)
        result = pd.Series(qh_values, index=pd.DatetimeIndex(qh_index), name="conso_kwh")
    else:
        # Pas plus fin que 15 min : simple agregation par sommation d'energie
        # (les kWh de sous-pas s'additionnent directement en kWh du quart d'heure).
        result = conso_kwh.resample("15min").sum()

    print(f"  /!\\ Conso reechantillonnee de {step_minutes} min vers 15 min "
          f"({'repartition uniforme' if step_minutes > 15 else 'agregation'}).")
    return result


def fill_missing_periods_by_calendar_symmetry(conso_kwh: pd.Series, step_minutes: int = 15) -> tuple:
    """
    Complete un releve de consommation incomplet -- que ce soit un trou interne
    (ex: une semaine ou un mois manquant au milieu de l'annee) ou une periode
    couverte plus courte qu'une annee complete (ex: un releve qui ne va que
    de janvier a septembre) -- en reconstituant une grille reguliere complete
    sur l'ANNEE CIVILE de la premiere date du releve (1er janvier -> 31
    decembre), au pas de temps `step_minutes`.

    Principe (symetrie calendaire) : pour chaque quart d'heure manquant, on
    prend la valeur du jour disponible le plus proche en distance calendaire
    CYCLIQUE (nombre de jours sur une annee de 365, sans distinction
    avant/apres), a la meme heure/minute -- ex. s'il manque la 2e moitie de
    l'automne, elle est completee avec les jours de la 1ere moitie d'automne
    les plus proches (ou, a defaut, les jours equivalents une saison plus
    loin), pas avec une valeur d'hiver ou d'ete qui n'aurait pas de sens.

    Meme regle qu'il s'agisse d'un trou en debut/fin de periode ou d'un trou
    interne : la grille cible est toujours l'annee civile complete. Si le
    releve depasse cette annee civile (ex: 14 mois, a cheval sur 2 annees),
    l'exces est tronque -- le moteur de simulation tourne toujours sur un
    horizon d'exactement 1 an.

    Retourne (serie_completee, n_points_tronques) : n_points_tronques est le
    nombre de pas de temps du releve original qui tombaient hors de l'annee
    civile retenue (0 si le releve ne depassait pas 12 mois).
    """
    if len(conso_kwh) < 2:
        return conso_kwh, 0

    # La grille cible est toujours l'annee civile complete de la premiere date
    # du releve (peu importe le mois de depart reel des donnees) : 1er janvier
    # 00:00 -> 31 decembre 23:45 (dernier quart d'heure de l'annee), au pas de
    # temps demande.
    year = conso_kwh.index[0].year
    full_index = pd.date_range(
        pd.Timestamp(year=year, month=1, day=1),
        pd.Timestamp(year=year, month=12, day=31, hour=23, minute=60 - step_minutes),
        freq=f"{step_minutes}min")

    in_year_mask = (conso_kwh.index >= full_index[0]) & (conso_kwh.index <= full_index[-1])
    n_truncated = int((~in_year_mask).sum())
    if n_truncated:
        conso_kwh = conso_kwh[in_year_mask]
        print(f"  /!\\ {n_truncated} pas de temps du releve tombent hors de l'annee civile "
              f"{year} (releve de plus de 12 mois) -- tronques, la simulation tourne sur "
              f"exactement 1 an ({year}).")

    missing_mask = ~full_index.isin(conso_kwh.index)
    n_missing = int(missing_mask.sum())
    if n_missing == 0:
        return conso_kwh, n_truncated

    # Index (heure, minute) -> jour le plus proche disponible (distance
    # calendaire cyclique) -- voir engine/calendar_utils.py, partage avec
    # l'alignement des prix Day-Ahead (engine/fournisseur_roi.py).
    calendar_index = NearestByCalendarSymmetry()
    for ts, val in conso_kwh.items():
        calendar_index.add(subkey=(ts.hour, ts.minute), month=ts.month, day=ts.day, value=val)
    calendar_index.finalize()

    fallback_mean = float(conso_kwh.mean())  # utilise si aucun jour comparable n'est disponible pour un (heure, minute) donne

    def _closest_value(ts) -> float:
        return calendar_index.closest(
            subkey=(ts.hour, ts.minute), month=ts.month, day=ts.day, fallback=fallback_mean)

    filled_values = []
    existing = conso_kwh.to_dict()
    n_filled = 0
    for ts in full_index:
        if ts in existing:
            filled_values.append(existing[ts])
        else:
            filled_values.append(_closest_value(ts))
            n_filled += 1

    print(f"  /!\\ {n_filled} pas de temps manquants dans le releve de consommation "
          f"({n_filled * step_minutes / 60:.0f} h au total) -- completes par symetrie "
          f"calendaire (jour disponible le plus proche, meme heure, meme rythme "
          f"hebdomadaire pas garanti).")

    filled_series = pd.Series(filled_values, index=full_index, name=conso_kwh.name or "conso_kwh")
    return filled_series, n_truncated


def conso_from_normalized_profile(profile: pd.Series, conso_annuelle_kwh: float) -> pd.Series:
    """
    Met a l'echelle un profil de consommation normalise (fraction de l'annee
    par quart d'heure, somme = 1.0 -- voir io_sources/extract_conso_profiles_from_excel.py)
    avec une consommation annuelle cible, pour obtenir une serie en kWh par
    quart d'heure directement utilisable par le reste du pipeline.
    """
    if conso_annuelle_kwh < 0:
        raise ValueError("La consommation annuelle cible doit etre positive ou nulle.")
    if len(profile) < 2:
        raise ValueError(
            f"Profil de consommation vide ou trop court ({len(profile)} point(s)) -- "
            f"verifie data/conso_profiles.pkl (le regenerer si besoin via "
            f"io_sources/extract_conso_profiles_from_excel.py)."
        )
    series = profile * conso_annuelle_kwh
    series.name = "conso_kwh"
    return series


def build_timeseries_from_sources(conso_series_or_csv, pv_hourly_profile_1kwc: pd.Series,
                                    kwc: float, unit: str = "kW", timestamp_col: str = None,
                                    value_col: str = None, injection_col: str = None) -> pd.DataFrame:
    """
    Construit le DataFrame (conso_kwh, pv_kwh, hour, is_weekend) attendu par
    le moteur de dispatch (fournisseur_roi.py), a partir :
    - d'une consommation client : soit un CSV (chemin/buffer, voir
      read_conso_csv), soit une pd.Series deja en kWh par quart d'heure
      (ex: un profil integre mis a l'echelle via conso_from_normalized_profile)
    - d'un profil PV PVGIS pour 1 kWc (voir pv_pvgis.fetch_pv_profile_1kwc)

    C'est le remplacement direct de extract_all() pour la nouvelle interface
    sans Excel : le reste du pipeline (run_fournisseur_model, etc.) n'a besoin
    d'aucune modification, il consomme ce DataFrame de la meme facon.

    Pour un CSV, timestamp_col / value_col / injection_col designent les colonnes
    a utiliser (auto-detectees si omises ; pas d'injection si injection_col est
    None). Si une injection est lue, elle est ajoutee au DataFrame (colonne
    "injection_mesuree_kwh", kWh par quart d'heure) a titre INFORMATIF : elle ne
    pilote pas la simulation.
    """
    from io_sources.pv_pvgis import expand_to_quarter_hour  # import local pour eviter une dependance circulaire au chargement du module

    n_truncated = 0
    injection_qh = None
    if isinstance(conso_series_or_csv, pd.Series):
        # Profil integre deja mis a l'echelle (conso_from_normalized_profile) :
        # pas de completion calendaire necessaire, on suppose l'annee deja complete.
        print("Consommation client : profil integre deja en kWh/quart d'heure.")
        conso_qh = align_conso_to_quarter_hour(conso_series_or_csv)
    else:
        # Cas CSV utilisateur : le releve peut etre partiel ou a un pas de temps
        # different de 15 min, d'ou l'alignement puis la completion par symetrie
        # calendaire pour obtenir une annee civile complete exploitable par le LP.
        print("Lecture de la consommation client (CSV)...")
        conso_raw, injection_raw = read_conso_injection_csv(
            conso_series_or_csv, timestamp_col=timestamp_col, value_col=value_col,
            injection_col=injection_col, unit=unit)
        conso_qh = align_conso_to_quarter_hour(conso_raw)
        conso_qh, n_truncated = fill_missing_periods_by_calendar_symmetry(conso_qh)
        if injection_raw is not None:
            # Meme traitement que la conso (grille 15 min puis annee civile complete),
            # pour que les deux series restent alignees pas a pas.
            injection_qh = align_conso_to_quarter_hour(injection_raw)
            injection_qh, _ = fill_missing_periods_by_calendar_symmetry(injection_qh)
            injection_qh = injection_qh.rename("injection_mesuree_kwh")

    print("Alignement du profil PV (PVGIS) sur la grille quart-horaire du client...")
    pv_qh = expand_to_quarter_hour(pv_hourly_profile_1kwc, kwc=kwc, target_index=conso_qh.index)

    df = pd.concat([conso_qh, pv_qh], axis=1).sort_index()
    df["conso_kwh"] = df["conso_kwh"].fillna(0.0)
    df["pv_kwh"] = df["pv_kwh"].fillna(0.0)
    if injection_qh is not None:
        df["injection_mesuree_kwh"] = injection_qh.reindex(df.index).fillna(0.0)
    df.index.name = "datetime"
    df["hour"] = df.index.hour
    df["is_weekend"] = df.index.dayofweek >= 5  # 5=samedi, 6=dimanche (convention pandas dayofweek)

    print(f"  -> Serie finale : {len(df)} points, conso={df['conso_kwh'].sum():,.0f} kWh/an, "
          f"pv={df['pv_kwh'].sum():,.0f} kWh/an")
    df.attrs["n_truncated_qh"] = n_truncated  # trace, pour affichage dans l'UI, du nombre de pas tronques lors de la completion calendaire
    return df


def extract_parameters(xlsm_path: str = XLSM_PATH) -> dict:
    """Raccourci pour ne recuperer que les parametres (sans les series
    temporelles) via le chemin legacy Excel `extract_all`. Relit tout le
    classeur (parametres + series) puis jette les series -- pas optimise,
    mais coherent avec `extract_all` par construction."""
    params, _ = extract_all(xlsm_path)
    return params


def extract_timeseries(xlsm_path: str = XLSM_PATH, kwc: float = None) -> pd.DataFrame:
    """Raccourci pour ne recuperer que le DataFrame de series temporelles
    (sans les parametres) via le chemin legacy Excel `extract_all`. Voir
    extract_parameters : meme remarque sur le cout de relecture complete."""
    _, df = extract_all(xlsm_path, kwc=kwc)
    return df
